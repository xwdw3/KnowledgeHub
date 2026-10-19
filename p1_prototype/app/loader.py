"""文档解析：txt / md / pdf / docx → 纯文本。"""
from __future__ import annotations

import re
from pathlib import Path

ALLOW_EXT = {".txt", ".md", ".markdown", ".pdf", ".docx"}


class ParseError(Exception):
    """文档解析失败。"""


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:          # 单页异常不应导致整篇失败
            pages.append("")
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:       # 表格内容也是知识，按行拼接
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def clean_text(raw: str) -> str:
    """清洗：统一换行、压缩空白、去除控制字符与重复空行。"""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ord(ch) >= 32)
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_file(file_path: str | Path) -> tuple[str, str]:
    """返回 (文件名, 清洗后的文本)。"""
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext not in ALLOW_EXT:
        raise ParseError(f"不支持的文件类型：{ext}，仅支持 {'/'.join(sorted(ALLOW_EXT))}")

    if ext in {".txt", ".md", ".markdown"}:
        raw = _read_text(path)
    elif ext == ".pdf":
        raw = _read_pdf(path)
    else:
        raw = _read_docx(path)

    text = clean_text(raw)
    if not text:
        raise ParseError("文档解析成功，但未提取到任何文本（可能是扫描件或加密文档）")
    return path.name, text
